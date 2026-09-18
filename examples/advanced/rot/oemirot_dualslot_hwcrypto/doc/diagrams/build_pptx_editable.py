#!/usr/bin/env python3
"""
Build the deck out of NATIVE PowerPoint shapes — every box, line and label is a
real object you can select, move, recolour and retype. No pictures.

    python3 build_pptx_editable.py        ->  oemirot_architecture.pptx

It replays the same drawing calls the SVG generator uses (generate_diagrams.py
in "rec" mode), so both decks always show the same content.

Fidelity note: PowerPoint lays text out with its own font metrics, so labels can
sit a pixel or two off where the SVG puts them. If you need pixel-exact slides
for printing, use build_pptx.py, which embeds the rendered images instead.
"""
import os
import re
import sys

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_PATTERN, MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate_diagrams as G
from build_pptx import SLIDES

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "oemirot_architecture.pptx")

UNIT = 12192000 / 1600.0          # EMU per SVG unit (16:9 slide, 1600x900 canvas)
PT_PER_UNIT = 0.6                 # 900 units == 540 pt


def E(v):
    return Emu(int(round(v * UNIT)))


def PTS(v):
    return Pt(round(v * PT_PER_UNIT, 1))


def rgb(c):
    return RGBColor.from_string(c.lstrip("#").upper())


def font_of(family):
    return "Consolas" if "Consolas" in (family or "") else "Segoe UI"


# ------------------------------------------------------------------ fills ---
def apply_fill(shape, spec):
    if spec is None or spec == "none":
        shape.fill.background()
        return
    if isinstance(spec, str) and spec.startswith("url(#hatch"):
        shape.fill.patterned()
        shape.fill.pattern = MSO_PATTERN.LIGHT_UPWARD_DIAGONAL
        if "Gray" in spec:
            shape.fill.fore_color.rgb = rgb(G.GRAY_BD)
            shape.fill.back_color.rgb = rgb(G.GRAY_LT)
        else:
            shape.fill.fore_color.rgb = rgb(G.AMBER)
            shape.fill.back_color.rgb = rgb(G.AMBER_LT)
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(spec)


def apply_line(shape, color, width, dash):
    if color is None:
        shape.line.fill.background()
        return
    shape.line.color.rgb = rgb(color)
    shape.line.width = E(width)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def add_arrowhead(shape):
    ln = shape.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"),
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)


# ------------------------------------------------------------- primitives ---
def draw_rect(slide, x, y, w, h, fill, stroke, rx, sw, dash):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rx and rx > 0 else MSO_SHAPE.RECTANGLE,
        E(x), E(y), E(w), E(h))
    shp.shadow.inherit = False
    if rx and rx > 0:
        try:
            shp.adjustments[0] = min(0.5, float(rx) / max(1.0, min(w, h)))
        except (IndexError, ValueError):
            pass
    apply_fill(shp, fill)
    apply_line(shp, stroke, sw, dash)
    shp.text_frame.word_wrap = False
    return shp


def draw_circle(slide, x, y, r, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x - r), E(y - r),
                                 E(2 * r), E(2 * r))
    shp.shadow.inherit = False
    apply_fill(shp, fill)
    shp.line.fill.background()
    return shp


def draw_text(slide, x, y, s, size, fill, anchor, weight, family):
    if not s:
        return None
    bold = weight in ("600", "700", "bold")
    # generous, so the box never forces a wrap; alignment does the positioning
    est = max(len(s) * size * (0.78 if bold else 0.72), 40)
    if anchor == "middle":
        left, align = x - est / 2.0, PP_ALIGN.CENTER
    elif anchor == "end":
        left, align = x - est, PP_ALIGN.RIGHT
    else:
        left, align = x, PP_ALIGN.LEFT
    top = y - size * 1.02
    box = slide.shapes.add_textbox(E(left), E(top), E(est), E(size * 1.75))
    tf = box.text_frame
    # wrap="none" + autofit makes PowerPoint shrink the box and re-centre it on
    # its own midpoint, which drags left-aligned labels to the right.
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = s
    run.font.size = PTS(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(fill)
    run.font.name = font_of(family)
    return box


def draw_line(slide, x1, y1, x2, y2, color, sw, dash, marker):
    con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                     E(x1), E(y1), E(x2), E(y2))
    con.line.color.rgb = rgb(color)
    con.line.width = E(sw)
    if dash:
        con.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if marker:
        add_arrowhead(con)
    return con


TOKEN = re.compile(r"[MLCZz]|-?\d+(?:\.\d+)?")


def path_points(d):
    """Flatten an SVG path (M/L/C only) into a polyline."""
    toks = TOKEN.findall(d)
    pts, cur, cmd, i = [], None, None, 0
    while i < len(toks):
        t = toks[i]
        if t in "MLCZz":
            cmd = t.upper()
            i += 1
            continue
        if cmd in ("M", "L"):
            x, y = float(toks[i]), float(toks[i + 1])
            i += 2
            pts.append((x, y))
            cur = (x, y)
        elif cmd == "C":
            x1, y1, x2, y2, x, y = (float(v) for v in toks[i:i + 6])
            i += 6
            p0 = cur or (x1, y1)
            for k in range(1, 15):
                u = k / 14.0
                bx = ((1 - u) ** 3 * p0[0] + 3 * (1 - u) ** 2 * u * x1 +
                      3 * (1 - u) * u ** 2 * x2 + u ** 3 * x)
                by = ((1 - u) ** 3 * p0[1] + 3 * (1 - u) ** 2 * u * y1 +
                      3 * (1 - u) * u ** 2 * y2 + u ** 3 * y)
                pts.append((bx, by))
            cur = (x, y)
        else:
            i += 1
    return pts


def draw_path(slide, d, stroke, sw, dash, marker):
    pts = path_points(d)
    for i in range(len(pts) - 1):
        last = (i == len(pts) - 2)
        draw_line(slide, pts[i][0], pts[i][1], pts[i + 1][0], pts[i + 1][1],
                  stroke, sw, dash, marker if (marker and last) else None)


def replay(slide, records):
    for rec in records:
        kind = rec[0]
        if kind == "rect":
            draw_rect(slide, *rec[1:])
        elif kind == "circle":
            draw_circle(slide, *rec[1:])
        elif kind == "text":
            draw_text(slide, *rec[1:])
        elif kind == "line":
            draw_line(slide, *rec[1:])
        elif kind == "path":
            draw_path(slide, *rec[1:])


# ------------------------------------------------------------------ build ---
def title_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                            prs.slide_width, prs.slide_height)
    bg.shadow.inherit = False
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(G.INK)
    bg.line.fill.background()
    draw_text(s, 80, 380, "Secure boot and firmware update", 52, G.WHITE,
              "start", "700", G.FONT)
    draw_text(s, 80, 440, "Architecture options for OEMiRoT on STM32C5", 30,
              G.CYAN, "start", "400", G.FONT)
    draw_text(s, 80, 486, "based on example_oemirot_dualslot_hwcrypto", 19,
              G.WHITE, "start", "400", G.FONT)


def main():
    G.set_mode("rec")
    for fn in G.DIAGRAMS:
        fn()
    G.set_mode("svg")

    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    title_slide(prs)

    missing = []
    for name, note in SLIDES:
        if name not in G.SLIDES_REC:
            missing.append(name)
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        replay(slide, G.SLIDES_REC[name])
        slide.notes_slide.notes_text_frame.text = note

    prs.save(OUT)
    n = len(prs.slides._sldIdLst)
    print(f"editable deck -> {OUT}  ({n} slides, native shapes)")
    if missing:
        print("  missing diagrams:", ", ".join(missing))


if __name__ == "__main__":
    main()
