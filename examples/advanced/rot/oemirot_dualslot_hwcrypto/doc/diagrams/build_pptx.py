#!/usr/bin/env python3
"""
Assemble the generated diagrams into a 16:9 PowerPoint deck with speaker notes.

Run generate_diagrams.py first, then:  python3 build_pptx.py
Output: oemirot_architecture.pptx  (one full-bleed diagram per slide)

The deck is a starting point — drop the slides into your own template, or take
the SVGs from svg/ and place them on your own layouts.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PNG = os.path.join(HERE, "png")
OUT = os.path.join(HERE, "oemirot_architecture.pptx")

NAVY = RGBColor(0x03, 0x23, 0x4B)
CYAN = RGBColor(0x3C, 0xB4, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDES = [
    ("00-glossary",
     "Vocabulary slide. Do not read it out - put it up, say it stays available, and come "
     "back to it whenever someone asks what a term means."),
    ("01-two-stage-boot",
     "The frame for everything else: an immutable first stage that checks the second one, "
     "then hides itself. Nothing the application does can reach back into it."),
    ("11-trust-model",
     "Signing and encryption answer different questions. Note the two weak points we own: "
     "a single key pair with no revocation path, and a decryption key sitting raw in flash."),
    ("02-flash-map",
     "Real offsets from flash_layout.h. The headline number is 464 KB usable out of 1 MB - "
     "that is the price of the second slot, and it drives most of the later decisions."),
    ("03-update-pipeline",
     "Key point: the bootloader never downloads anything. All transport code lives in the "
     "replaceable application, so the immutable part stays small and offline."),
    ("04-boot-decision",
     "This runs on every reset, not just on update. Three independent rejection paths, and "
     "the primary slot is re-verified even when no update is pending."),
    ("05-overwrite-sequence",
     "What the board does today. Walk the five states, then land on the red box: a signed "
     "image that crashes takes the update path down with it."),
    ("05b-why-install-request",
     "Answers the question the previous slide always provokes: why a separate 'request "
     "install' step at all. The short version: the two stages cannot talk, so the only "
     "channel is a mark in flash - and that mark also proves the download completed."),
    ("06-swap-sequence",
     "The same flow plus one confirmation flag. Worth stressing: the confirm code already "
     "exists in our application, compiled out by OVERWRITE_ONLY."),
    ("07-bank-swap",
     "Attractive on paper - instant install. Be honest about the open question: SWAP_BANK "
     "moves the bootloader too, and nothing in our repo implements this yet."),
    ("08-architecture-variants",
     "The four arrangements side by side. The closing question is the one that actually "
     "picks a column; everything else follows from it."),
    ("09-comparison",
     "Detail view of the same trade-offs. The status row separates what we have from what "
     "we would have to build."),
    ("10-failure-modes",
     "The argument in one table. Four failures are handled identically - only the last row "
     "separates overwrite from swap."),
    ("12-decision-tree",
     "Proposed order for the decision. Questions 5 and 6 are irreversible once devices are "
     "provisioned, so they need an answer before anything ships."),
    ("13-faq-images",
     "FAQ. The magic trailer is a completion flag, not a trust signal - and the app writes "
     "the image encrypted because it has no access to the key. Keep for questions."),
    ("14-faq-data",
     "FAQ. Calibration data must sit outside the slots; mirroring inside them does not work. "
     "Note the cost: reserving N bytes removes N from each slot."),
    ("15-faq-swap-and-speed",
     "FAQ. Two corrections worth stressing: swap needs ONE bootloader and no mirroring, and "
     "encryption costs the running application nothing."),
    ("16-open-questions",
     "Closing slide. Do not try to answer these here - the goal is to leave with an owner and "
     "a date against each. The key strategy and the lock-down level are the two that cannot "
     "be revisited after provisioning."),
]


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(11.5), Inches(2.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Secure boot and firmware update"
    r.font.size, r.font.bold, r.font.color.rgb = Pt(44), True, WHITE
    r.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Architecture options for OEMiRoT on STM32C5"
    r2.font.size, r2.font.color.rgb, r2.font.name = Pt(26), CYAN, "Segoe UI"
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    r3.text = "based on example_oemirot_dualslot_hwcrypto"
    r3.font.size, r3.font.color.rgb, r3.font.name = Pt(16), WHITE, "Segoe UI"

    for name, note in SLIDES:
        img = os.path.join(PNG, name + ".png")
        if not os.path.exists(img):
            print(f"missing {img} - run generate_diagrams.py first")
            continue
        sl = prs.slides.add_slide(blank)
        sl.shapes.add_picture(img, 0, 0, width=prs.slide_width,
                              height=prs.slide_height)
        sl.notes_slide.notes_text_frame.text = note

    prs.save(OUT)
    print(f"deck -> {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
